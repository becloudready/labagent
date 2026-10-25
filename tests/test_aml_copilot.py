"""AML pack, part 4: mode=copilot. Builds real Word and Excel files and opens them again the way the software would."""
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from labagent.packs.aml.copilot import guidance as G
from labagent.packs.aml.copilot import package as P
from labagent.packs.aml.copilot.exercises import SECTIONS, build_exercises
from labagent.packs.aml.copilot.select import select_cases
from labagent.packs.aml.gen import get
from labagent.packs.aml.pack import AmlLab
from labagent.packs.base import load_state, run_pack
from tests.test_aml_corpus import ECFR, JUSTICE

PAGE = "<html><nav>menu</nav><main><h1>Indicators</h1>" + "".join(f"<p>Indicator {i}: a client behaviour that may point to laundering.</p>" for i in range(30)) + "</main><footer>f</footer></html>"


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("LABAGENT_HOME", str(tmp_path / "home"))
    return tmp_path


def run(home, action="launch", **params):
    base = {"mode": "copilot", "size": "tiny", "cases": 12, "out": str(home / "out")}
    base.update(params)
    sv = {"allow_tiny": True, "regs_http": lambda u: ECFR if "ECFR" in u else JUSTICE, "regs_cache": home / "regs", "guidance_http": lambda u: PAGE.encode()}
    return run_pack(AmlLab(), action, "cohort1", base, services=sv)


def text_of(path):
    d = Document(str(path))
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """One launch shared by the tests that only read the result."""
    home = tmp_path_factory.mktemp("copilot")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("LABAGENT_HOME", str(home / "home"))
        env = run(home)
    assert env.status == "ok", env.error
    return home, env


@pytest.fixture(scope="module")
def tables():
    return get("builtin").generate(seed=42, size="small", jurisdiction="ca").tables


# ---- the launch ----------------------------------------------------------------------------------------------------------------

def test_launch_builds_a_learner_folder_an_instructor_folder_and_two_zips(built):
    home, env = built
    out = home / "out"
    learner, instr = out / "cohort1", out / "cohort1-instructor"
    assert sorted(p.name for p in learner.iterdir()) == ["Cases", "Data", "Exercise Guide.docx", "Policies", "Regulations"]
    assert sorted(p.name for p in instr.iterdir()) == ["Exercise Guide with Answers.docx", "Instructor Answer Key.docx"]
    assert len(list((learner / "Cases").glob("Case [0-9]*.xlsx"))) == 12 == len(list((learner / "Cases").glob("Case [0-9]*.docx")))
    assert len(list((learner / "Data" / "Transactions by month").glob("*.xlsx"))) == 12
    assert {p.name for p in (learner / "Policies").iterdir()} == {"AML Policy Manual (Canada).docx", "Alert Handling Procedure (Canada).docx", "Guide to Writing the Report Narrative (Canada).docx"}
    o = env.summary["outputs"]
    assert o["mode"] == "copilot" and o["package"]["checks_passed"] and o["package"]["exercises"] == 13 and Path(o["learner_zip"]).exists()


def test_the_learner_zip_has_no_instructor_file_and_the_instructor_zip_has_the_answers(built):
    home, _ = built
    lz = zipfile.ZipFile(home / "out" / "cohort1.zip").namelist()
    iz = zipfile.ZipFile(home / "out" / "cohort1-instructor.zip").namelist()
    assert lz and all(n.startswith("cohort1/") for n in lz) and not any(re.search(r"instructor|answer", n, re.I) for n in lz)
    assert sorted(iz) == ["cohort1-instructor/Exercise Guide with Answers.docx", "cohort1-instructor/Instructor Answer Key.docx"]


def test_every_workbook_holds_real_excel_tables_with_clean_headers(built):
    home, _ = built
    books = list((home / "out" / "cohort1").rglob("*.xlsx"))
    assert len(books) == 12 + 1 + 1 + 12 + 4          # cases, case list, all transactions, months, customers/accounts/alerts/screening
    for b in books:
        wb = load_workbook(b)
        for ws in wb.worksheets:
            assert ws.tables, (b.name, ws.title)
            heads = [c.value for c in ws[1]]
            assert all(heads) and len(set(heads)) == len(heads)
            assert not ws.merged_cells.ranges


def test_case_workbooks_have_the_four_tabs_and_no_answer(built, tables):
    home, _ = built
    wb = load_workbook(next((home / "out" / "cohort1" / "Cases").glob("Case 01 - *.xlsx")))
    assert wb.sheetnames == ["Summary", "KYC", "Transactions", "Alerts"]
    flat = " ".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value)
    assert not any(g["reason"] in flat for g in tables["ground_truth"]) and "escalate" not in flat.lower()
    tx = wb["Transactions"]
    assert [c.value for c in tx[1]][:3] == ["Date and time", "ID", "From account"] and tx.max_row > 2


def test_learner_files_are_free_of_technical_terms_and_answers(built):
    home, _ = built
    for p in (home / "out" / "cohort1").rglob("*.docx"):
        t = text_of(p)
        assert not P.JARGON.search(t), (p.name, P.JARGON.search(t).group(0))
        if p.parent.name != "Regulations":
            assert "—" not in t and "Expected answer" not in t, p.name
    assert P.verify_package(home / "out" / "cohort1", home / "out" / "cohort1-instructor", []) == []


def test_exercise_guide_has_thirteen_exercises_and_one_catch_per_section(built):
    home, _ = built
    learner = text_of(home / "out" / "cohort1" / "Exercise Guide.docx")
    key = text_of(home / "out" / "cohort1-instructor" / "Exercise Guide with Answers.docx")
    assert len(re.findall(r"^[A-D]\d\. ", learner, re.M)) == 13 and "Expected answer" not in learner and key.count("Expected answer") == 13
    assert len(re.findall(r"Catch the AI's mistake", learner)) == 4
    for sentence in re.split(r"(?<=[.!?])\s+", learner):
        assert len(sentence.split()) <= 40, sentence[:80]
    for needle in ("AutoSave", "OneDrive or SharePoint", "five files"):
        assert needle in learner


def test_answer_key_matches_the_ground_truth(built):
    home, _ = built
    from labagent.packs.aml.schema import read_csv
    t = read_csv(home / "out" / ".work" / "tables")
    key = text_of(home / "out" / "cohort1-instructor" / "Instructor Answer Key.docx")
    cases = select_cases(t, 12)
    for c in cases:
        assert f"{c['number']:02d} | {c['alert']['id']} | {c['truth']['label']}" in key.replace("\n", " | ") or c["truth"]["reason"] in key
        assert c["truth"]["reason"] in key


def test_regulation_excerpts_and_fintrac_guidance_are_written_with_their_notices(built):
    home, env = built
    names = [p.name for p in (home / "out" / "cohort1" / "Regulations").iterdir()]
    assert any("Act (excerpt)" in n for n in names) and any(n.startswith("FINTRAC") for n in names)
    act = next(p for p in (home / "out" / "cohort1" / "Regulations").iterdir() if "Act (excerpt)" in p.name)
    assert "not the official version" in text_of(act)
    f = text_of(next(p for p in (home / "out" / "cohort1" / "Regulations").iterdir() if p.name.startswith("FINTRAC")))
    assert "Source: https://fintrac-canafe.canada.ca/" in f and "Indicator 7" in f and "menu" not in f
    assert env.summary["outputs"]["package"]["fintrac_guidance"]


def test_guidance_failure_is_reported_and_does_not_stop_the_launch(home):
    def down(u):
        raise OSError("blocked")
    env = run_pack(AmlLab(), "launch", "cohort1", {"mode": "copilot", "size": "tiny", "cases": 12, "out": str(home / "out")},
                   services={"allow_tiny": True, "regs_http": lambda u: ECFR if "ECFR" in u else JUSTICE, "regs_cache": home / "regs", "guidance_http": down})
    assert env.status == "ok" and len(env.summary["outputs"]["package"]["guidance_notes"]) == 2
    assert not any(p.name.startswith("FINTRAC") for p in (home / "out" / "cohort1" / "Regulations").iterdir())


def test_guidance_and_regulations_can_be_skipped(home):
    env = run(home, guidance="skip", regulations="skip")
    assert env.status == "ok"
    assert [p.name for p in (home / "out" / "cohort1" / "Regulations").iterdir()] == ["Read me.docx"]


def test_us_cohort_gets_us_rules_and_no_fintrac_guidance(home):
    env = run(home, jurisdiction="us")
    assert env.status == "ok"
    names = [p.name for p in (home / "out" / "cohort1" / "Regulations").iterdir()]
    assert names == ["US rules 31 CFR Chapter X (excerpt).docx"] and (home / "out" / "cohort1" / "Policies" / "AML Policy Manual (United States).docx").exists()


def test_launch_resumes_and_destroy_removes_only_its_own_files(home):
    run(home)
    assert run(home).summary["skipped"] == 3
    (home / "out" / "cohort1" / "My notes.docx").write_text("mine")
    env = run(home, action="destroy")
    assert env.status == "ok"
    left = [p.relative_to(home / "out").as_posix() for p in (home / "out").rglob("*") if p.is_file()]
    assert left == ["cohort1/My notes.docx"]
    assert not (home / "home" / "packs" / "aml" / "cohort1.json").exists()


def test_cohort_name_sets_the_folder_and_is_made_safe(home):
    env = run(home, cohort="Spring 2026/../x")
    assert env.status == "ok" and (home / "out" / "Spring 2026..x").is_dir()


# ---- exercises -----------------------------------------------------------------------------------------------------------------

def test_expected_answers_are_computed_from_the_data(tables):
    cases = select_cases(tables, 20)
    ex = {e["id"]: e for e in build_exercises(tables, cases, "ca", "c")}
    assert len(ex) == 13 and {e["section"] for e in ex.values()} == set(SECTIONS)
    et = [float(x["amount"]) for x in tables["transactions"] if x["channel"] == "e-transfer" and x["currency"] == "CAD"]
    assert f"{len(et):,}" in ex["D4"]["expected"][0] and f"{sum(et) / len(et):,.2f}" in ex["D4"]["expected"][1]
    assert f"{len(et) + 1000:,}" in ex["D4"]["draft"] and f"{sum(et) / len(et):,.2f}" not in ex["D4"]["draft"]
    sanc = next(c for c in cases if c["truth"]["pattern"] == "sanctions_match" and c["truth"]["label"] == "escalate")
    assert f"Case {sanc['number']:02d}" in ex["C3"]["expected"][0]
    top_name = ex["D2"]["expected"][0].split(". ", 1)[1].split(":")[0]
    assert top_name in {c["name"] for c in tables["customers"]}


def test_planted_mistakes_really_differ_from_the_files(tables):
    cases = select_cases(tables, 20)
    ca = {e["id"]: e for e in build_exercises(tables, cases, "ca", "c")}
    us = {e["id"]: e for e in build_exercises(tables, cases, "us", "c")}
    assert "more than $10,000" in ca["B2"]["draft"] and "$10,000 or more" in ca["B2"]["expected"][0]          # Canada says 10,000 or more
    assert "$10,000 or more" in us["B2"]["draft"] and "more than $10,000" in us["B2"]["expected"][0]          # the US says more than 10,000
    assert "three years" in ca["A3"]["draft"] and "every month" in ca["A3"]["draft"]


def test_case_selection_covers_every_kind_and_numbers_follow_alert_order(tables):
    cases = select_cases(tables, 20)
    assert [c["number"] for c in cases] == list(range(1, 21)) and [c["alert"]["id"] for c in cases] == sorted(c["alert"]["id"] for c in cases)
    kinds = {c["truth"]["pattern"] for c in cases}
    assert {"fan_in", "cycle", "stack", "disguised_payroll", "sanctions_match", "possible_cycle", "none"} <= kinds
    assert select_cases(tables, 20) == cases


def test_guidance_parser_keeps_the_main_text_only():
    text = G.html_to_text(PAGE)
    assert text.startswith("Indicators") and "menu" not in text and text.count("Indicator ") >= 30
    out, problems = G.fetch_guidance(lambda u: b"<html><main><p>short</p></main></html>")
    assert out == {} and len(problems) == 2
