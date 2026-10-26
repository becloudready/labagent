"""Build the Copilot folder: Policies, Regulations, Cases, Data and the exercise guide for learners, and a separate instructor folder
with the answers. The learner folder never contains an answer."""
from __future__ import annotations

import re
import zipfile
from datetime import date
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from docx import Document
from openpyxl import load_workbook

from ..corpus import documents as D
from ..schema import INSTITUTION, RULES, Tables
from .exercises import SECTIONS, build_exercises, case_facts
from .office import write_docx, write_plain_docx, write_xlsx
from .select import select_cases

EXCERPTS = {
    "us": [("31 CFR", ["1010.311", "1010.306", "1010.313", "1010.314", "1010.230", "1010.430", "1020.220", "1020.320"])],
    "ca": [("PCMLTFA", ["6", "6.1", "7", "8", "9.3", "9.6"]), ("PCMLTF STR Regulations", ["9"]), ("PCMLTF Regulations", ["10", "132", "148"])],
}
EXCERPT_TITLE = {"31 CFR": "US rules: 31 CFR Chapter X (excerpt)", "PCMLTFA": "Canada: Proceeds of Crime (Money Laundering) and Terrorist Financing Act (excerpt)",
                 "PCMLTF STR Regulations": "Canada: Suspicious Transaction Reporting Regulations (excerpt)", "PCMLTF Regulations": "Canada: Proceeds of Crime (Money Laundering) and Terrorist Financing Regulations (excerpt)"}
PLAIN = {"fan_out": "One account sends money to many accounts", "fan_in": "Many accounts send money to one account",
         "gather_scatter": "Money is gathered into one account, then split out", "scatter_gather": "Money is split across accounts, then gathered again",
         "cycle": "Money moves around a loop and comes back", "random": "Money passes through a chain of accounts", "bipartite": "One group of accounts pays another group",
         "stack": "Layers of accounts pay each other", "disguised_payroll": "Illegal money paid out as payroll", "disguised_supplies": "Illegal money paid out as supplier invoices",
         "sanctions_match": "The customer name matches a sanctions list entry", "possible_fan_in": "Looks like many-to-one, but small gifts",
         "possible_cash_structuring": "Cash deposits close to the limit at a cash business", "possible_cycle": "A loan and its repayment",
         "possible_rapid_in_out": "A payment passed straight on to a supplier", "none": "Ordinary activity with a documented reason"}
JARGON = re.compile(r"\b(RAG|embedding|embeddings|vector|vectors|Databricks|Lakebase|pgvector|chunk|chunks|LLM|API|JSON|CSV|Python|labagent|ground[_ ]truth)\b", re.I)
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿]")


def _name(cohort: str) -> str:
    return re.sub(r"[^A-Za-z0-9 _.-]", "", cohort).strip() or "cohort"


def excerpt_blocks(chunks: List[Dict[str, Any]], jurisdiction: str) -> Dict[str, List[Dict[str, Any]]]:
    """{file title: blocks} for the sections listed in EXCERPTS. Each file carries the source's notice."""
    by: Dict[str, Dict[str, List[str]]] = {}
    meta: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for c in chunks:
        m = c["meta"]
        by.setdefault(m["instrument"], {}).setdefault(m["section"], []).append(c["content"])
        meta[(m["instrument"], m["section"])] = m
    out: Dict[str, List[Dict[str, Any]]] = {}
    for j in (["ca", "us"] if jurisdiction == "both" else [jurisdiction]):
        for inst, sections in EXCERPTS[j]:
            blocks: List[Dict[str, Any]] = []
            notice = ""
            for sec in sections:
                parts = by.get(inst, {}).get(sec)
                if not parts:
                    continue
                m = meta[(inst, sec)]
                notice = m["notice"]
                blocks += [{"h1": f"{m['citation']}  {m['heading']}".strip()}, {"p": "\n\n".join(parts)}, {"p": f"Source: {m['url']} (version of {m['as_of']})"}]
            if blocks:
                out[EXCERPT_TITLE[inst]] = [{"quote": notice}] + blocks
    return out


def case_sheets(t: Tables, ix: D.CaseIndex, case: Dict[str, Any]) -> List[Dict[str, Any]]:
    a = case["alert"]
    c = ix.customer(a)
    rows, _ = ix.rows(a)
    return [
        {"title": "Summary", "name": "Summary", "columns": ["Item", "Value"], "rows": [
            ["Case number", f"{case['number']:02d}"], ["Alert ID", a["id"]], ["Date raised", a["date"]], ["Rule triggered", a["rule"]], ["What the rule looks for", RULES[a["rule"]]],
            ["Alert amount", a["amount"]], ["Risk score", str(a["risk_score"])], ["Account", a["account"]], ["Customer", c["name"]], ["Customer ID", c["id"]]]},
        {"title": "KYC", "name": "KYC", "columns": ["Field", "Value"], "rows": [
            ["Name", c["name"]], ["Type", c["type"]], ["Occupation or industry", c["occupation"]], ["Country", c["country"]], ["Risk rating", c["risk_rating"]],
            ["Expected monthly activity", str(c["expected_monthly_activity"])], ["Customer since", c["onboarding_date"]], ["Profile", ix.kyc[c["id"]]]]},
        {"title": "Transactions", "name": "Transactions", "columns": ["Date and time", "ID", "From account", "To account", "Amount", "Currency", "Channel", "Description"],
         "rows": [[x["timestamp"], x["id"], x["from_account"], x["to_account"], x["amount"], x["currency"], x["channel"], x["description"]] for x in rows]},
        {"title": "Alerts", "name": "Alerts", "columns": ["Alert ID", "Date", "Rule", "Amount", "Risk score"],
         "rows": [[o["id"], o["date"], o["rule"], o["amount"], str(o["risk_score"])] for o in ix.alerts_by_account.get(a["account"], [])]},
    ]


def data_files(t: Tables) -> Dict[str, List[Dict[str, Any]]]:
    tx = sorted(t["transactions"], key=lambda x: (x["timestamp"], x["id"]))
    cols = ["Date and time", "ID", "From account", "To account", "Amount", "Currency", "Channel", "Description"]
    row = lambda x: [x["timestamp"], x["id"], x["from_account"], x["to_account"], x["amount"], x["currency"], x["channel"], x["description"]]
    files: Dict[str, List[Dict[str, Any]]] = {"Transactions": [{"title": "Transactions", "name": "Transactions", "columns": cols, "rows": [row(x) for x in tx]}]}
    months: Dict[str, List[Any]] = {}
    for x in tx:
        months.setdefault(x["timestamp"][:7], []).append(row(x))
    for m, rows in sorted(months.items()):
        files[f"Transactions by month/Transactions {m}"] = [{"title": "Transactions", "name": "Transactions", "columns": cols, "rows": rows}]
    kyc = {k["customer_id"]: k["profile"] for k in t["kyc_profiles"]}
    files["Customers"] = [{"title": "Customers", "name": "Customers", "columns": ["ID", "Name", "Type", "Occupation or industry", "Country", "Risk rating", "Expected monthly activity", "Customer since", "Profile"],
                           "rows": [[c["id"], c["name"], c["type"], c["occupation"], c["country"], c["risk_rating"], str(c["expected_monthly_activity"]), c["onboarding_date"], kyc[c["id"]]] for c in t["customers"]]}]
    files["Accounts"] = [{"title": "Accounts", "name": "Accounts", "columns": ["ID", "Customer ID", "Type", "Opened", "Branch"],
                          "rows": [[a["id"], a["customer_id"], a["type"], a["open_date"], a["branch"]] for a in t["accounts"]]}]
    files["Alerts"] = [{"title": "Alerts", "name": "Alerts", "columns": ["ID", "Rule", "What the rule looks for", "Account", "Date", "Amount", "Risk score"],
                        "rows": [[a["id"], a["rule"], RULES[a["rule"]], a["account"], a["date"], a["amount"], str(a["risk_score"])] for a in t["alerts"]]}]
    files["Sanctions screening results"] = [{"title": "Results", "name": "Screening", "columns": ["ID", "Customer ID", "Customer name", "List", "List entry name", "Match type", "Score", "List date", "Note"],
                                             "rows": [[h["id"], h["customer_id"], h["customer_name"], h["list"], h["sdn_name"], h["match_type"], h["score"], h["list_date"], h["note"]] for h in t["sanctions_hits"]]}]
    return files


def guide_blocks(exercises: List[Dict[str, Any]], answers: bool, cohort: str, jurisdiction: str) -> List[Dict[str, Any]]:
    country = {"ca": "Canada", "us": "the United States", "both": "Canada"}[jurisdiction]
    b: List[Dict[str, Any]] = [
        {"p": f"{INSTITUTION} is an invented credit union. Every name, account and transfer in this folder is made up. The rules and laws are real ({country})." + (" This copy has the expected answers. Do not share it with learners." if answers else "")},
        {"h1": "Before you start"},
        {"numbered": ["Upload the whole folder to OneDrive or SharePoint. Copilot only works on files stored there.", "Open each Excel file from OneDrive or SharePoint. Check that AutoSave is on.",
                      "In Excel, click inside the table before you ask Copilot a question.", "Copilot can look at up to five files at a time. Pick your files before you ask.",
                      "If Copilot cannot handle the big Transactions file, use one month from Data > Transactions by month."]},
        {"h1": "How to check an AI answer"},
        {"bullets": ["Find the source. Every fact must point to a row, a cell or a sentence in a file.", "Do the sums yourself. Add up the amounts and count the rows.",
                     "Look hardest at numbers, dates and small words such as more than, at least and within."]},
        {"p": f"Cohort: {cohort}. There are {len(exercises)} exercises in four parts. Each part ends with one exercise where you find mistakes in an AI draft."},
    ]
    for sec in SECTIONS:
        group = [e for e in exercises if e["section"] == sec]
        if not group:
            continue
        b.append({"h1": f"Part: {sec}"})
        for e in group:
            b += [{"h2": f"{e['id']}. {e['title']}"}, {"numbered": e["steps"]}]
            if e.get("draft"):
                b.append({"quote": e["draft"]})
            b.append({"p": "Check yourself: " + " ".join(e["check"])})
            if answers:
                b += [{"p": "Expected answer:"}, {"bullets": e["expected"]}]
    return b


def answer_key_blocks(cases: List[Dict[str, Any]], t: Tables) -> List[Dict[str, Any]]:
    b: List[Dict[str, Any]] = [{"p": "For instructors only. Decisions come from how the data was built, not from a real investigation. Use them to guide the discussion."},
                               {"table": ["Case | Alert | Decision | What happened"] + [
                                   f"{c['number']:02d} | {c['alert']['id']} | {c['truth']['label']} | {PLAIN.get(c['truth']['pattern'], c['truth']['pattern'])}" for c in cases]}]
    for c in cases:
        f = case_facts(t, c)
        b += [{"h2": f"Case {c['number']:02d} ({c['alert']['id']}): {c['truth']['label']}"}, {"p": c["truth"]["reason"]},
              {"p": f"Customer: {f['customer']['name']}. Key transfers: {', '.join(x['id'] for x in f['keys'][:12])}{' and more' if len(f['keys']) > 12 else ''}."}]
    return b


def zip_folder(src: Path, dest: Path) -> Path:
    """A zip with the folder as its top level, in sorted order with fixed times."""
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(x for x in src.rglob("*") if x.is_file()):
            info = zipfile.ZipInfo(f"{src.name}/{p.relative_to(src).as_posix()}", (2026, 1, 1, 0, 0, 0))
            info.compress_type, info.external_attr = zipfile.ZIP_DEFLATED, 0o644 << 16
            z.writestr(info, p.read_bytes())
    return dest


def build_package(t: Tables, jurisdiction: str, cohort: str, out: Path, n_cases: int, reg_chunks: List[Dict[str, Any]],
                  guidance: Optional[Dict[str, str]] = None, with_data_months: bool = True) -> Dict[str, Any]:
    cohort = _name(cohort)
    learner, instr = out / cohort, out / f"{cohort}-instructor"
    country = lambda j: {"ca": "Canada", "us": "United States"}[j]
    created: List[str] = []

    for doc in D.policy_documents(jurisdiction):
        base = {"policy": "AML Policy Manual", "procedure": "Alert Handling Procedure", "guide": "Guide to Writing the Report Narrative"}[doc.kind]
        created.append(str(write_docx(doc, learner / "Policies" / f"{base} ({country(doc.jurisdiction)}).docx").relative_to(out)))

    ex_files = excerpt_blocks(reg_chunks, jurisdiction)
    for title, blocks in ex_files.items():
        created.append(str(write_plain_docx(title, blocks, learner / "Regulations" / f"{re.sub(r'[:/]', '', title)}.docx").relative_to(out)))
    for title, text in (guidance or {}).items():
        blocks = [{"p": para} for para in text.split("\n\n") if para.strip()]
        created.append(str(write_plain_docx(title, blocks, learner / "Regulations" / f"{title}.docx").relative_to(out)))
    if not ex_files and not guidance:
        created.append(str(write_plain_docx("Regulations were not included", [{"p": "This cohort folder was built without the regulation text. Ask your instructor to build it again with the regulation text included."}],
                                            learner / "Regulations" / "Read me.docx").relative_to(out)))

    cases = select_cases(t, n_cases)
    ix = D.CaseIndex(t)
    case_docs = {c.meta["alert_id"]: c for c in D.case_documents(t, jurisdiction, ix)}
    for c in cases:
        stem = f"Case {c['number']:02d} - {c['alert']['id']}"
        created.append(str(write_xlsx(learner / "Cases" / f"{stem}.xlsx", case_sheets(t, ix, c)).relative_to(out)))
        d = case_docs[c["alert"]["id"]]
        d.title = f"Case {c['number']:02d} file ({c['alert']['id']})"
        created.append(str(write_docx(d, learner / "Cases" / f"{stem}.docx").relative_to(out)))
    created.append(str(write_xlsx(learner / "Cases" / "Case list.xlsx", [{"title": "Cases", "name": "Cases", "columns": ["Case", "Alert ID", "Date raised", "Rule", "Account", "Alert amount"],
                                  "rows": [[f"{c['number']:02d}", c["alert"]["id"], c["alert"]["date"], c["alert"]["rule"], c["alert"]["account"], c["alert"]["amount"]] for c in cases]}]).relative_to(out)))

    for name, sheets in data_files(t).items():
        if name.startswith("Transactions by month/") and not with_data_months:
            continue
        created.append(str(write_xlsx(learner / "Data" / f"{name}.xlsx", sheets).relative_to(out)))

    exercises = build_exercises(t, cases, jurisdiction, cohort)
    created.append(str(write_plain_docx(f"Copilot exercises: {INSTITUTION}", guide_blocks(exercises, False, cohort, jurisdiction), learner / "Exercise Guide.docx").relative_to(out)))
    created.append(str(write_plain_docx(f"Copilot exercises with answers: {INSTITUTION}", guide_blocks(exercises, True, cohort, jurisdiction), instr / "Exercise Guide with Answers.docx").relative_to(out)))
    created.append(str(write_plain_docx(f"Answer key: {INSTITUTION}", answer_key_blocks(cases, t), instr / "Instructor Answer Key.docx").relative_to(out)))

    lz = zip_folder(learner, out / f"{cohort}.zip")
    iz = zip_folder(instr, out / f"{cohort}-instructor.zip")
    created += [str(lz.relative_to(out)), str(iz.relative_to(out))]
    return {"learner_dir": str(learner), "instructor_dir": str(instr), "learner_zip": str(lz), "instructor_zip": str(iz), "cases": len(cases),
            "exercises": len(exercises), "files": len(created), "created": created, "excerpts": len(ex_files), "guidance": sorted(guidance or {})}


def _docx_text(path: Path) -> str:
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        parts += [c.text for r in tb.rows for c in r.cells]
    return "\n".join(parts)


def verify_package(learner: Path, instructor: Path, reasons: List[str]) -> List[str]:
    """Reopen every file as a person's software would. Returns problems (empty is good)."""
    problems: List[str] = []
    for p in sorted(learner.rglob("*")):
        if not p.is_file():
            continue
        rel = p.relative_to(learner).as_posix()
        if re.search(r"instructor|answer", p.name, re.I) and p.name != "Exercise Guide.docx":
            problems.append(f"{rel}: an instructor file is in the learner folder")
        if p.suffix == ".xlsx":
            try:
                wb = load_workbook(p, read_only=False)
            except Exception as e:
                problems.append(f"{rel}: cannot open ({type(e).__name__})")
                continue
            for ws in wb.worksheets:
                if not ws.tables:
                    problems.append(f"{rel}: sheet {ws.title} has no Excel table")
                heads = [c.value for c in ws[1]]
                if len(set(heads)) != len(heads) or any(h in (None, "") for h in heads):
                    problems.append(f"{rel}: sheet {ws.title} has empty or repeated headers")
            if "/Data/" not in "/" + rel and len(wb.worksheets) < 5:
                text = " ".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows(max_row=60) for c in r if isinstance(c.value, str))
                if JARGON.search(text) or any(r in text for r in reasons):
                    problems.append(f"{rel}: technical term or answer text")
        elif p.suffix == ".docx":
            try:
                text = _docx_text(p)
            except Exception as e:
                problems.append(f"{rel}: cannot open ({type(e).__name__})")
                continue
            if rel.startswith("Regulations/") and p.name.startswith("FINTRAC"):
                continue
            if JARGON.search(text):
                problems.append(f"{rel}: technical term {JARGON.search(text).group(0)!r}")
            if ("\u2014" in text and not rel.startswith("Regulations/")) or EMOJI.search(text):      # quoted law keeps its own punctuation
                problems.append(f"{rel}: em dash or emoji")
            if any(r in text for r in reasons):
                problems.append(f"{rel}: contains answer text")
    for need in ("Exercise Guide.docx",):
        if not (learner / need).exists():
            problems.append(f"missing {need}")
    if not (instructor / "Instructor Answer Key.docx").exists():
        problems.append("missing Instructor Answer Key.docx")
    return problems
